"""文档解析：把 PDF / DOCX 转成带页码映射的纯文本。

页码映射的作用：切分条款后要记录 char_start / char_end，
前端用它在原文里做高亮定位，同时也要知道这一条在第几页。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}

_MULTI_BLANK = re.compile(r"\n{3,}")
_TRAILING_SPACE = re.compile(r"[ \t]+\n")


@dataclass
class ParsedDocument:
    """解析结果：全文 + 页码区间（page_no, start, end）。"""

    text: str
    page_offsets: list[tuple[int, int, int]] = field(default_factory=list)

    @property
    def page_count(self) -> int:
        return len(self.page_offsets) or 1

    def page_no_at(self, char_offset: int) -> int:
        """根据字符偏移反查页码。"""
        for page_no, start, end in self.page_offsets:
            if start <= char_offset < end:
                return page_no
        return self.page_offsets[-1][0] if self.page_offsets else 1


def _normalize(text: str) -> str:
    """统一换行、去掉多余空行与行尾空白。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u3000", " ")  # 全角空格
    text = _TRAILING_SPACE.sub("\n", text)
    text = _MULTI_BLANK.sub("\n\n", text)
    return text.strip()


def parse_document(path: Path) -> ParsedDocument:
    """按扩展名选择解析器。"""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix == ".docx":
        return _parse_docx(path)
    raise ValueError(f"不支持的文件类型：{suffix}")


def _parse_pdf(path: Path) -> ParsedDocument:
    """逐页抽取文本，同时记录每页在全文中的字符区间。"""
    reader = PdfReader(str(path))
    parts: list[str] = []
    offsets: list[tuple[int, int, int]] = []
    cursor = 0

    for page_no, page in enumerate(reader.pages, start=1):
        page_text = _normalize(page.extract_text() or "")
        if cursor > 0:
            parts.append("\n")
            cursor += 1
        start = cursor
        parts.append(page_text)
        cursor += len(page_text)
        offsets.append((page_no, start, cursor))

    return ParsedDocument(text="".join(parts), page_offsets=offsets)


def _parse_docx(path: Path) -> ParsedDocument:
    """抽取段落与表格文本。

    注意：Word 的分页是渲染时决定的，python-docx 拿不到真实页码，
    所以 DOCX 统一记为第 1 页；定位仍可依赖 char_start / char_end。
    """
    doc = Document(str(path))
    blocks: list[str] = [p.text for p in doc.paragraphs]

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            blocks.append(" | ".join(cells))

    text = _normalize("\n".join(blocks))
    return ParsedDocument(text=text, page_offsets=[(1, 0, len(text))])
